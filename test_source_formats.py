"""Real format fixtures and end-to-end checks; contains synthetic questions only."""
import io
import json
import re
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.sax.saxutils import escape

import pandas as pd
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from streamlit.testing.v1 import AppTest

import app
import question_importer as qi
import source_formats as sf
from test_regressions import make_mcq_32_source

QUESTIONS = ['Which number is even?', 'Which shape has three sides?']
OPTIONS = [['Two', 'Three', 'Five', 'Seven'], ['Circle', 'Triangle', 'Square', 'Pentagon']]
ANSWERS = ['A', 'B']
TEXT = '\n\n'.join('\n'.join([f'{i+1}. {q}', *[f'{chr(65+j)}. {v}' for j,v in enumerate(OPTIONS[i])], f'ANSWER: {ANSWERS[i]}']) for i,q in enumerate(QUESTIONS))
JSON_ROWS = [{'number': i+1, 'question': q, 'options': OPTIONS[i], 'correct': ANSWERS[i]} for i,q in enumerate(QUESTIONS)]
TABLE_ROWS = [['No','Question','Option 1','Option 2','Option 3','Option 4','Correct'], *[[i+1,q,*OPTIONS[i],ANSWERS[i]] for i,q in enumerate(QUESTIONS)]]


def uploaded(name, data):
    result = io.BytesIO(data)
    result.name = name
    return result


def docx_bytes():
    doc = Document()
    for line in TEXT.splitlines():
        doc.add_paragraph(line)
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    for row in TABLE_ROWS:
        sheet.append(row)
    output = io.BytesIO(); book.save(output)
    return output.getvalue()


def change_content_type(data, main_part, mimetype):
    from lxml import etree
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            blob = src.read(name)
            if name == '[Content_Types].xml':
                root = etree.fromstring(blob)
                for item in root:
                    if item.get('PartName') == main_part:
                        item.set('ContentType',mimetype)
                blob = etree.tostring(root)
            dst.writestr(name,blob)
    return output.getvalue()


def pdf_bytes(text=TEXT, graphics=False):
    writer = PdfWriter(); page = writer.add_blank_page(width=595,height=842)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    commands = ['BT /F1 12 Tf 50 810 Td']
    for line in text.splitlines():
        literal = line.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
        commands.append(f'0 -18 Td ({literal}) Tj')
    commands.append('ET')
    if graphics:
        commands.append('10 10 20 20 re S')
    stream = DecodedStreamObject();stream.set_data('\n'.join(commands).encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO();writer.write(out);return out.getvalue()


def ods_bytes():
    rows = ''.join('<table:table-row>'+''.join('<table:table-cell office:value-type="string"><text:p>'+escape(str(value))+'</text:p></table:table-cell>' for value in row)+'</table:table-row>' for row in TABLE_ROWS)
    content = f'''<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:spreadsheet><table:table table:name="Questions">{rows}</table:table></office:spreadsheet></office:body></office:document-content>'''
    out = io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        z.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet')
        z.writestr('content.xml',content)
        z.writestr('META-INF/manifest.xml','''<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>''')
    return out.getvalue()


def xml_bytes():
    nodes=[]
    for i,q in enumerate(QUESTIONS):
        choices=''.join(f'<answer fraction="{100 if j==i else 0}"><text>{escape(v)}</text></answer>' for j,v in enumerate(OPTIONS[i]))
        nodes.append(f'<question type="multichoice"><questiontext format="html"><text>{escape(q)}</text></questiontext><single>true</single>{choices}</question>')
    return ('<?xml version="1.0"?><quiz>'+''.join(nodes)+'</quiz>').encode()


def convert_fixture(data, target, directory):
    exe=shutil.which('libreoffice') or shutil.which('soffice')
    if not exe:
        raise RuntimeError('Install libreoffice-writer to run the format compatibility tests.')
    root=Path(directory); source=root/'fixture.docx';source.write_bytes(data)
    out=root/target;out.mkdir(exist_ok=True)
    filters={'doc':'doc:MS Word 97','dot':'dot:MS Word 97 Vorlage','rtf':'rtf:Rich Text Format','odt':'odt:writer8','ott':'ott:writer8_template','fodt':'fodt:OpenDocument Text Flat XML'}
    p=subprocess.run([exe,'-env:UserInstallation='+(root/'fixture-profile').as_uri(),'--headless','--convert-to',filters[target],'--outdir',str(out),str(source)],capture_output=True,timeout=90)
    path=out/('fixture.'+target)
    if p.returncode or not path.exists():
        raise AssertionError(f'Could not produce a real {target} fixture: '+p.stderr.decode(errors='replace'))
    return path.read_bytes()


def make_fixtures(directory):
    import csv
    doc=docx_bytes();xlsx=workbook_bytes()
    fixtures={'docx':doc,'xlsx':xlsx,'txt':TEXT.encode(),'ods':ods_bytes(),'pdf':pdf_bytes(),
              'json':json.dumps({'questions':JSON_ROWS}).encode(),'jsonl':'\n'.join(json.dumps(row) for row in JSON_ROWS).encode(),'xml':xml_bytes()}
    for ext in ['csv','tsv']:
        out=io.StringIO();writer=csv.writer(out,delimiter=',' if ext=='csv' else '\t');writer.writerows(TABLE_ROWS);fixtures[ext]=out.getvalue().encode()
    for ext in ['md','markdown']:
        fixtures[ext]='\n'.join('**'+line+'**' if re.match(r'\d+\.',line) else line for line in TEXT.splitlines()).encode()
    for ext in ['html','htm']:
        fixtures[ext]=('<!doctype html><html><head><title>Synthetic quiz</title></head><body>'+''.join('<p>'+escape(line)+'</p>' for line in TEXT.splitlines())+'</body></html>').encode()
    wordtypes={'docm':'application/vnd.ms-word.document.macroEnabled.main+xml','dotx':'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml','dotm':'application/vnd.ms-word.template.macroEnabledTemplate.main+xml'}
    for ext,mime in wordtypes.items():
        fixtures[ext]=change_content_type(doc,'/word/document.xml',mime)
    exceltypes={'xlsm':'application/vnd.ms-excel.sheet.macroEnabled.main+xml','xltx':'application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml','xltm':'application/vnd.ms-excel.template.macroEnabled.main+xml'}
    for ext,mime in exceltypes.items():
        fixtures[ext]=change_content_type(xlsx,'/xl/workbook.xml',mime)
    for ext in ['doc','dot','rtf','odt','ott','fodt']:
        fixtures[ext]=convert_fixture(doc,ext,directory)
    return fixtures


class FormatCompatibilityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix='mcq_format_tests_')
        cls.fixtures=make_fixtures(cls.temp.name)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_every_new_extension_has_a_real_fixture(self):
        self.assertEqual(set(sf.SUPPORTED_EXTENSIONS)-{'xls'},set(self.fixtures))
        self.assertTrue(self.fixtures['doc'].startswith(bytes.fromhex('D0CF11E0A1B11AE1')))
        self.assertTrue(self.fixtures['dot'].startswith(bytes.fromhex('D0CF11E0A1B11AE1')))
        self.assertTrue(self.fixtures['rtf'].startswith(b'{\\rtf'))

    def test_all_formats_preserve_questions_options_answers_and_export(self):
        for ext,data in self.fixtures.items():
            with self.subTest(format=ext):
                frame,report,images=qi.import_all_sources([uploaded('sample.'+ext,data)])
                self.assertFalse(app.has_blocking_errors(report),report.to_dict('records'))
                self.assertEqual(frame.Question.tolist(),QUESTIONS)
                self.assertEqual(frame.Correct.tolist(),ANSWERS)
                self.assertEqual(frame[qi.OPTION_COLUMNS[:4]].values.tolist(),OPTIONS)
                reviewed,validation=qi.validate_reviewed_questions(frame)
                self.assertFalse(app.has_blocking_errors(validation),validation.to_dict('records'))
                versions=app.build_versions(reviewed,3,'Letters',True,True)
                for questions in versions.values():
                    for question in questions:
                        index=QUESTIONS.index(question['Question'])
                        correct=next(o['text'] for o in question['Options'] if o['letter']==question['Correct'])
                        self.assertEqual(correct,OPTIONS[index][index])
                package=app.create_zip_output_from_versions(versions,reviewed,report,'Compatibility check',True,image_map=images)
                with zipfile.ZipFile(package) as archive:
                    self.assertIsNone(archive.testzip())
                    for label in ['A','B','C']:
                        self.assertIn('Exam_'+label+'.docx',archive.namelist())
                        document=Document(io.BytesIO(archive.read('Exam_'+label+'.docx')))
                        body='\n'.join(p.text for p in document.paragraphs)
                        self.assertTrue(all(q in body for q in QUESTIONS))

    def test_every_new_format_in_streamlit_upload_review_and_create(self):
        for ext,data in self.fixtures.items():
            if ext in {'docx','xlsx','txt'}:
                continue
            with self.subTest(format=ext):
                at=AppTest.from_file('app.py',default_timeout=90).run()
                at.get('file_uploader')[0].upload('sample.'+ext,data,'application/octet-stream').run()
                self.assertEqual(len(at.exception),0)
                self.assertEqual(at.session_state['combined_df'].Question.tolist(),QUESTIONS)
                confirmations=[c for c in at.checkbox if c.label.startswith('I confirm that I reviewed')]
                if confirmations:
                    confirmations[0].check().run()
                next(b for b in at.button if b.label=='Create exam package').click().run()
                self.assertEqual(len(at.exception),0)
                self.assertTrue(at.session_state['preview_versions'])
                self.assertIn('4 — Download',[h.value for h in at.header])
                self.assertTrue(any('Download exam package' in b.label for b in at.get('download_button')))

    def test_converted_document_images_survive(self):
        for ext in ['doc','odt','rtf']:
            with self.subTest(format=ext):
                data=convert_fixture(make_mcq_32_source(),ext,self.temp.name)
                frame,issues,images=qi.import_all_sources([uploaded('figure.'+ext,data)])
                self.assertEqual(len(frame),1,issues.to_dict('records'))
                self.assertTrue(images,'The converted question lost its embedded image.')
                self.assertTrue(frame.Needs_Review.all())
                self.assertTrue(issues.loc[issues.Severity == 'WARNING', 'Issue'].str.contains('image size').any())

    def test_clean_legacy_conversions_are_informational(self):
        for ext in sf.CONVERTED_DOCUMENTS:
            with self.subTest(format=ext):
                frame, report, _ = qi.import_all_sources([uploaded('sample' + ext, self.fixtures[ext[1:]])])
                self.assertEqual(frame.Correct.tolist(), ANSWERS)
                self.assertFalse(frame.Needs_Review.any())
                self.assertEqual(frame.Confidence.tolist(), ['High', 'High'])
                self.assertEqual(report.Severity.tolist(), ['INFO'])
                self.assertTrue(frame.Import_Warnings.eq('').all())

    def test_converted_highlighted_and_missing_answers_keep_their_confidence(self):
        document = Document()
        for index, question in enumerate(QUESTIONS):
            document.add_paragraph(f'{index + 1}. {question}')
            for option_index, option in enumerate(OPTIONS[index]):
                run = document.add_paragraph().add_run(f'{chr(65 + option_index)}. {option}')
                if index == 0 and option_index == 0:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        output = io.BytesIO()
        document.save(output)
        data = convert_fixture(output.getvalue(), 'doc', self.temp.name)
        frame, report, _ = qi.import_all_sources([uploaded('highlighted.doc', data)])
        self.assertEqual(frame.Correct.tolist(), ['A', ''])
        self.assertEqual(frame.Needs_Review.tolist(), [False, True])
        self.assertEqual(frame.iloc[0].Answer_Evidence, 'Highlighted answer option')
        self.assertEqual(report.Severity.tolist(), ['INFO'])
        _, validation = qi.validate_reviewed_questions(frame)
        self.assertTrue(app.has_blocking_errors(validation))

    def test_clean_legacy_conversion_has_no_false_ui_warning(self):
        at = AppTest.from_file('app.py', default_timeout=60).run()
        at.get('file_uploader')[0].upload('sample.doc', self.fixtures['doc'], 'application/msword').run()
        self.assertFalse(at.exception)
        self.assertTrue(any('Imported from DOC: 2 questions read.' in item.value for item in at.info))
        self.assertFalse(at.warning)
        self.assertFalse(at.error)
        self.assertFalse([c for c in at.checkbox if c.label.startswith('I confirm that I reviewed')])
        self.assertTrue(any('No validation issues detected.' in item.value for item in at.success))
        next(b for b in at.button if b.label == 'Create exam package').click().run()
        self.assertFalse(at.exception)
        self.assertTrue(at.session_state['preview_versions'])

    def test_native_docx_bypasses_conversion(self):
        with patch('source_formats.convert_document', side_effect=AssertionError('DOCX must not be converted')):
            frame, report, _ = qi.import_all_sources([uploaded('sample.docx', self.fixtures['docx'])])
        self.assertEqual(frame.Correct.tolist(), ANSWERS)
        self.assertTrue(report.empty)

    def test_import_is_reused_in_session_and_same_size_replacement_is_detected(self):
        at=AppTest.from_file('app.py',default_timeout=60).run()
        with patch('question_importer.import_all_sources',wraps=qi.import_all_sources) as importer:
            at.get('file_uploader')[0].upload('bank.csv',self.fixtures['csv'],'text/csv').run()
            count=importer.call_count
            next(n for n in at.number_input if n.label=='Number of exam versions').set_value(2).run()
            self.assertEqual(importer.call_count,count)
            replacement=self.fixtures['csv'].replace(b'Two',b'Six')
            at.get('file_uploader')[0].upload('bank.csv',replacement,'text/csv').run()
            self.assertGreater(importer.call_count,count)
            self.assertEqual(at.session_state['combined_df'].iloc[0]['Option 1'],'Six')


class FormatEdgeCaseTests(unittest.TestCase):
    def assert_rejected(self,name,data,message=None):
        frame,report,_=qi.import_all_sources([uploaded(name,data)])
        self.assertTrue(app.has_blocking_errors(report),report.to_dict('records'))
        if message:
            self.assertTrue(report.Issue.str.contains(message,case=False,regex=False).any(),report.to_dict('records'))
        return frame,report

    def test_csv_semicolon_quotes_multiline_and_utf16(self):
        import csv
        rows=[TABLE_ROWS[0],[1,'Which option includes a comma, quote " and\nline break?','Yes, indeed','No','Maybe','Never','A']]
        for delimiter,encoding in [(';','utf-8-sig'),(',','utf-16')]:
            out=io.StringIO();csv.writer(out,delimiter=delimiter).writerows(rows)
            frame,report,_=qi.import_all_sources([uploaded('bank.csv',out.getvalue().encode(encoding))])
            self.assertFalse(app.has_blocking_errors(report))
            self.assertEqual(frame.iloc[0]['Question'],rows[1][1])
            self.assertEqual(frame.iloc[0]['Option 1'],'Yes, indeed')

    def test_utf16_greek_and_numbered_same_file_key(self):
        text='32. Ποιο είναι σωστό;\nA. Γάτα\nB. Σκύλος\n\nAnswer Key\n32. B\n'
        frame,report,_=qi.import_all_sources([uploaded('greek.txt',text.encode('utf-16'))])
        self.assertEqual(frame.iloc[0]['Question'],'Ποιο είναι σωστό;')
        self.assertEqual(frame.Correct.tolist(),['B'])

    def test_markdown_table_and_static_html_table(self):
        markdown='| '+' | '.join(TABLE_ROWS[0])+' |\n| '+' | '.join(['---']*len(TABLE_ROWS[0]))+' |\n'
        markdown+='\n'.join('| '+' | '.join(map(str,row))+' |' for row in TABLE_ROWS[1:])
        html='<table>'+''.join('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in TABLE_ROWS)+'</table>'
        for name,text in [('table.md',markdown),('table.html',html)]:
            frame,report,_=qi.import_all_sources([uploaded(name,text.encode())])
            self.assertEqual(frame.Question.tolist(),QUESTIONS,report.to_dict('records'))
            self.assertEqual(frame.Correct.tolist(),ANSWERS)

    def test_broken_file_does_not_hide_valid_bank(self):
        frame,report,_=qi.import_all_sources([uploaded('broken.json',b'{bad'),uploaded('valid.txt',TEXT.encode())])
        self.assertEqual(len(frame),2)
        self.assertTrue(app.has_blocking_errors(report))

    def test_missing_correct_answer_is_blocked_at_review(self):
        row={'question':'Which number is even?','options':['Two','Three'],'correct':''}
        frame,_,_=qi.import_all_sources([uploaded('bank.json',json.dumps([row]).encode())])
        self.assertTrue(frame.Needs_Review.all())
        _,report=qi.validate_reviewed_questions(frame)
        self.assertTrue(app.has_blocking_errors(report))

    def test_unsupported_structures_do_not_silently_drop_content(self):
        for name,data,message in [
            ('empty.csv',b'','empty'),
            ('bad.csv',b'Question,Option 1,Option 2,Correct\nQ,A,B,C,extra','more fields'),
            ('bad.json',json.dumps([{'question':'Q','options':['a']*7,'correct':'A'}]).encode(),'2–6'),
            ('bad.json',json.dumps([{'question':'Q','options':['a','b'],'correct':['A','B']}]).encode(),'single'),
            ('bad.xml',b'<quiz><question type="essay"/></quiz>','not supported'),
            ('bad.xml',b'<!DOCTYPE quiz [<!ENTITY x SYSTEM "file:///etc/passwd">]><quiz/>','DTD'),
            ('image.html',b'<p>1. Which figure?</p><img src="https://example.invalid/image.png">','images'),
            ('image.md',b'1. Which figure?\n![image](https://example.invalid/image.png)','images'),
            ('scanned.pdf',pdf_bytes('',graphics=True),'no readable text'),
            ('figure.pdf',pdf_bytes(graphics=True),'graphics'),
            ('fake.docm',b'not a zip','container'),
        ]:
            with self.subTest(name=name,message=message):
                self.assert_rejected(name,data,message)

    def test_password_protected_pdf(self):
        writer=PdfWriter();writer.add_blank_page(100,100);writer.encrypt('password')
        out=io.BytesIO();writer.write(out)
        self.assert_rejected('locked.pdf',out.getvalue(),'Password-protected')

    def test_wrapped_pdf_and_html_preserve_stems_and_options(self):
        text='32. Which option has\na wrapped description?\nA. First choice with\na second line\nB. Other choice\nANSWER: A'
        html=''.join('<p>'+escape(line)+'</p>' for line in text.splitlines())
        for name,data in [('wrapped.pdf',pdf_bytes(text)),('wrapped.html',html.encode()),('wrapped.md',text.encode())]:
            with self.subTest(name=name):
                frame,report,_=qi.import_all_sources([uploaded(name,data)])
                self.assertFalse(app.has_blocking_errors(report),report.to_dict('records'))
                self.assertEqual(frame.Question.tolist(),['Which option has a wrapped description?'])
                self.assertEqual(frame.iloc[0]['Option 1'],'First choice with a second line')
                self.assertEqual(frame.Correct.tolist(),['A'])
                self.assertTrue(frame.Needs_Review.all())

    def test_missing_and_excess_options_are_rejected(self):
        seven='1. Which option?\n'+'\n'.join(f'{chr(65+i)}. Choice {i+1}' for i in range(7))+'\nANSWER: A'
        doc=Document()
        for line in seven.splitlines():
            doc.add_paragraph(line)
        out=io.BytesIO();doc.save(out)
        cases=[('seven.docx',out.getvalue(),'six'),('seven.txt',seven.encode(),'six'),
               ('seven.pdf',pdf_bytes(seven),'six'),
               ('gap.md',b'1. Which choice?\nA. First\nC. Third\nANSWER: A','missing'),
               ('gap.csv',b'Question,Option 1,Option 3,Correct\nQ,First,Third,A','missing'),
               ('extra.csv',b'Question,Option A,Option B,Option G,Correct\nQ,First,Second,Seventh,A','six')]
        for name,data,message in cases:
            with self.subTest(name=name):
                self.assert_rejected(name,data,message)

    def test_mixed_html_layout_and_ods_images_are_rejected(self):
        self.assert_rejected('mixed.md',b'1. Outside question?\nA. Yes\nB. No\n| Question | Option 1 | Option 2 | Correct |\n| --- | --- | --- | --- |\n| Inside? | Yes | No | A |','mixes')
        html='<p>1. Outside question?</p><p>A. Yes</p><p>B. No</p><table>'+''.join('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in TABLE_ROWS)+'</table>'
        self.assert_rejected('mixed.html',html.encode(),'mixes')
        out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(ods_bytes())) as src,zipfile.ZipFile(out,'w') as dst:
            for name in src.namelist():
                data=src.read(name)
                if name=='content.xml':
                    data=data.replace(b'<text:p>No',b'<image/><text:p>No')
                dst.writestr(name,data)
        self.assert_rejected('image.ods',out.getvalue(),'images')

    def test_json_lettered_options_and_moodle_truefalse(self):
        row={'question':'Which number is even?','options':{'A':'Two','B':'Three'},'correct':1}
        xml=b'<quiz><question type="category"/><question type="truefalse"><questiontext><text>Two is even.</text></questiontext><answer fraction="100"><text>True</text></answer><answer fraction="0"><text>False</text></answer></question></quiz>'
        for name,data in [('bank.json',json.dumps([row]).encode()),('bank.xml',xml)]:
            frame,report,_=qi.import_all_sources([uploaded(name,data)])
            self.assertFalse(app.has_blocking_errors(report),report.to_dict('records'))
            self.assertEqual(frame.Correct.tolist(),['A'])

    def test_missing_converter_and_timeout_are_actionable(self):
        with patch('source_formats.shutil.which',return_value=None):
            self.assert_rejected('old.doc',b'synthetic','Save the file as DOCX')
        with patch('source_formats.shutil.which',return_value='/usr/bin/soffice'),patch('source_formats.subprocess.Popen') as popen,patch('source_formats.os.killpg'):
            process=popen.return_value;process.pid=123;process.wait.side_effect=[subprocess.TimeoutExpired('soffice',45),0]
            self.assert_rejected('old.doc',b'synthetic','too long')


if __name__=='__main__':
    unittest.main()
