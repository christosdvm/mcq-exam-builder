import io
import unittest
import zipfile
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw

from question_importer import (
    WordDrawingAsset,
    insert_preserved_word_drawing,
    read_docx_questions,
    validation_issue_location,
)


def make_mcq_32_source() -> bytes:
    source_image = Image.new("RGB", (400, 300), "white")
    drawing = ImageDraw.Draw(source_image)
    drawing.rectangle((0, 0, 99, 199), fill=(220, 20, 20))
    drawing.rectangle((100, 0, 299, 199), fill=(20, 180, 20))
    drawing.rectangle((300, 0, 399, 199), fill=(20, 20, 220))
    drawing.rectangle((0, 200, 399, 299), fill=(240, 200, 20))
    image_bytes = io.BytesIO()
    source_image.save(image_bytes, format="PNG")
    image_bytes.seek(0)

    document = Document()
    question = document.add_paragraph()
    question_run = question.add_run(
        "32. This is a longitudinal histologic picture of a peripheral nerve of an animal "
        "(haematoxylin & eosin). Which structure is indicated by the yellow arrows?"
    )
    question_run.bold = True

    image_paragraph = document.add_paragraph()
    picture_run = image_paragraph.add_run()
    picture_run.add_picture(image_bytes, width=Inches(4))
    blip = picture_run._r.xpath(".//a:blip")[0]
    source_rect = OxmlElement("a:srcRect")
    source_rect.set("l", "25000")
    source_rect.set("r", "25000")
    source_rect.set("b", "33333")
    blip_fill = blip.getparent()
    blip_fill.insert(blip_fill.index(blip) + 1, source_rect)

    for label, text in [
        ("A", "The synaptic cleft"),
        ("B", "The dendrites"),
        ("C", "The Schwann cell nucleus"),
        ("D", "The nodes of Ranvier"),
    ]:
        option = document.add_paragraph()
        option_run = option.add_run(f"{label}. {text}")
        if label == "D":
            option_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_grouped_figure_source() -> bytes:
    """Create a synthetic Word figure with a raster and a numbered overlay."""
    image = Image.new("RGB", (240, 160), (228, 238, 247))
    drawing = ImageDraw.Draw(image)
    drawing.ellipse((35, 25, 205, 135), fill=(94, 129, 172))
    image_output = io.BytesIO()
    image.save(image_output, format="PNG")

    document = Document()
    document.add_paragraph("1. Which synthetic structure is labelled?")
    image_paragraph = document.add_paragraph()
    picture_run = image_paragraph.add_run()
    picture_run.add_picture(io.BytesIO(image_output.getvalue()), width=Inches(2))

    graphic_data = picture_run._r.xpath(".//a:graphicData")[0]
    picture = deepcopy(graphic_data.xpath("./pic:pic")[0])
    for child in list(graphic_data):
        graphic_data.remove(child)
    graphic_data.set("uri", "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup")
    group = parse_xml(
        """
        <wpg:wgp
          xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
          xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
          xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <wpg:cNvGrpSpPr/>
          <wpg:grpSpPr>
            <a:xfrm>
              <a:off x="0" y="0"/>
              <a:ext cx="1828800" cy="1219200"/>
              <a:chOff x="0" y="0"/>
              <a:chExt cx="1828800" cy="1219200"/>
            </a:xfrm>
          </wpg:grpSpPr>
          <wps:wsp>
            <wps:cNvPr id="2" name="Synthetic label"/>
            <wps:cNvSpPr/>
            <wps:spPr>
              <a:xfrm><a:off x="1250000" y="180000"/><a:ext cx="260000" cy="220000"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
              <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
              <a:ln><a:solidFill><a:srgbClr val="111827"/></a:solidFill></a:ln>
            </wps:spPr>
            <wps:txbx><w:txbxContent><w:p><w:r><w:t>SYNTHETIC LABEL 1</w:t></w:r></w:p></w:txbxContent></wps:txbx>
            <wps:bodyPr/>
          </wps:wsp>
        </wpg:wgp>
        """
    )
    group.insert(2, picture)
    graphic_data.append(group)

    for label in "ABCD":
        run = document.add_paragraph().add_run(f"{label}. Synthetic option {label}")
        if label == "A":
            run.bold = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_inline_answer_source() -> bytes:
    document = Document()
    for number, correct in [(1, "C"), (2, "B")]:
        document.add_paragraph(
            f"Question {number} — Which synthetic choice is correct for item {number}?"
        )
        for label in "ABCD":
            document.add_paragraph(f"{label}. Synthetic option {label}")
        document.add_paragraph(f"Correct answer: {correct}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_bold_answer_source(partial: bool = False) -> bytes:
    document = Document()
    document.add_paragraph("Which synthetic option is correct?")
    for label in "ABCD":
        paragraph = document.add_paragraph()
        if label == "B" and partial:
            paragraph.add_run("B. ")
            bold_run = paragraph.add_run("Synthetic correct option")
            bold_run.bold = True
        else:
            run = paragraph.add_run(f"{label}. Synthetic option {label}")
            if label == "B":
                run.bold = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _set_list_metadata(paragraph, num_id: int, level: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    list_level = OxmlElement("w:ilvl")
    list_level.set(qn("w:val"), str(level))
    numbering_id = OxmlElement("w:numId")
    numbering_id.set(qn("w:val"), str(num_id))
    numbering.append(list_level)
    numbering.append(numbering_id)
    properties.append(numbering)


def make_multilevel_unlabelled_source() -> bytes:
    document = Document()
    questions = [
        (
            "What type of placenta does the mare possess?",
            ["Cotyledonary", "Zonary", "Discoid", "Diffuse"],
            "D",
        ),
        (
            "A maiden mare has multiple small follicles but no dominant follicle. What is the diagnosis?",
            ["Ovulation failure", "Pregnancy", "Transitional phase", "Endometritis"],
            "C",
        ),
        (
            "The most predictive sign of ovulation in the mare is?",
            ["Small follicle", "Soft irregular follicle wall", "No uterine edema", "Tight cervix"],
            "B",
        ),
        (
            "The best time following insemination for the first pregnancy ultrasound is?",
            ["Day 10", "Day 16", "Day 30", "Day 60"],
            "B",
        ),
    ]
    for stem, options, correct in questions:
        question = document.add_paragraph(stem)
        _set_list_metadata(question, num_id=1, level=0)
        for index, option_text in enumerate(options):
            option = document.add_paragraph()
            run = option.add_run(option_text)
            if chr(ord("A") + index) == correct:
                run.bold = True
            _set_list_metadata(option, num_id=1, level=1)

    heading = document.add_paragraph()
    heading.add_run("Answer Key").bold = True
    for answer in "DCBB":
        paragraph = document.add_paragraph()
        paragraph.add_run(answer).bold = True
        _set_list_metadata(paragraph, num_id=5, level=0)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _add_four_option_question(document, number: int, missing_first_label: bool = False) -> None:
    document.add_paragraph(f"{number}. Which synthetic answer is correct for item {number}?")
    option_lines = []
    for label in "ABCD":
        prefix = "" if missing_first_label and label == "A" else f"{label}. "
        option_lines.append(f"{prefix}Synthetic option {label}")
    document.add_paragraph("\n".join(option_lines))


def make_numbered_answer_key_source() -> bytes:
    document = Document()
    for number in range(1, 7):
        _add_four_option_question(document, number, missing_first_label=number == 3)
    document.add_paragraph("answer key:")
    document.add_paragraph("1. C\n2. A\n3. D\n4. A\n5. C\n6. A")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pipe_answer_key_source() -> bytes:
    document = Document()
    for number in range(1, 4):
        _add_four_option_question(document, number)
    document.add_paragraph("Answer Key: 1. C | 2. D | 3. B")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_correct_answers_source_with_invalid_key() -> bytes:
    document = Document()
    for number in range(1, 3):
        _add_four_option_question(document, number)
    document.add_paragraph("Correct Answers:")
    document.add_paragraph("1. D")
    document.add_paragraph("2. F")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class GroupedWordFigureRegressionTests(unittest.TestCase):
    def test_group_overlay_and_original_extent_survive_import_and_export(self):
        frame, issues, image_map = read_docx_questions(
            make_grouped_figure_source(), "Synthetic_Grouped_Figure.docx"
        )

        self.assertEqual(len(frame), 1)
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])
        self.assertEqual(len(image_map), 1)
        asset = next(iter(image_map.values()))
        self.assertIsInstance(asset, WordDrawingAsset)
        self.assertIn(b"wpg:wgp", asset.drawing_xml)
        self.assertIn(b"SYNTHETIC LABEL 1", asset.drawing_xml)

        output_document = Document()
        run = output_document.add_paragraph().add_run()
        self.assertTrue(insert_preserved_word_drawing(output_document, run, asset))
        output = io.BytesIO()
        output_document.save(output)

        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
            document_xml = archive.read("word/document.xml")
            media_files = [
                name for name in archive.namelist() if name.startswith("word/media/")
            ]
        self.assertIn(b"wpg:wgp", document_xml)
        self.assertIn(b"SYNTHETIC LABEL 1", document_xml)
        self.assertIn(b'<wp:extent cx="1828800" cy="1219200"', document_xml)
        self.assertIn(b"<wp:inline", document_xml)
        self.assertNotIn(b"<wp:anchor", document_xml)
        self.assertEqual(len(media_files), 1)


class Mcq32CropRegressionTests(unittest.TestCase):
    def test_mcq_32_preserves_word_crop_through_round_trip(self):
        frame, issues, image_map = read_docx_questions(
            make_mcq_32_source(), "MCQ_32_Cropped_Image.docx"
        )

        self.assertEqual(len(frame), 1)
        self.assertEqual(str(frame.iloc[0]["No"]), "32")
        self.assertEqual(frame.iloc[0]["Correct"], "D")
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])

        image_names = frame.iloc[0]["Image_Files"].split("|")
        self.assertEqual(len(image_names), 1)
        cropped_blob = image_map[image_names[0].lower()]
        with Image.open(io.BytesIO(cropped_blob)) as cropped:
            self.assertEqual(cropped.size, (200, 200))
            self.assertEqual(cropped.getpixel((10, 100)), (20, 180, 20))
            self.assertEqual(cropped.getpixel((190, 100)), (20, 180, 20))

        exported = Document()
        exported.add_paragraph().add_run().add_picture(
            io.BytesIO(cropped_blob), width=Inches(4)
        )
        exported_bytes = io.BytesIO()
        exported.save(exported_bytes)
        with zipfile.ZipFile(io.BytesIO(exported_bytes.getvalue())) as archive:
            media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
            with Image.open(io.BytesIO(archive.read(media_name))) as round_trip:
                self.assertEqual(round_trip.size, (200, 200))

    def test_warning_location_identifies_mcq_32(self):
        location = validation_issue_location(
            {
                "Source file": "MCQ_32_Cropped_Image.docx",
                "Row": "32",
                "Internal_ID": "MCQ_32_Cropped_Image-32",
            }
        )
        self.assertEqual(
            location,
            "MCQ_32_Cropped_Image.docx · question/row 32 · ID MCQ_32_Cropped_Image-32",
        )


class LecturerWordFormatRegressionTests(unittest.TestCase):
    def test_numbered_multiline_answer_key_and_missing_first_option_label(self):
        frame, issues, _ = read_docx_questions(
            make_numbered_answer_key_source(), "Numbered_Key.docx"
        )

        self.assertEqual(len(frame), 6)
        self.assertEqual(frame["Correct"].tolist(), ["C", "A", "D", "A", "C", "A"])
        self.assertEqual(frame["Option_Count"].tolist(), [4, 4, 4, 4, 4, 4])
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])
        self.assertFalse([item for item in issues if "Source labels" in item["Issue"]])

    def test_pipe_separated_answer_key_on_one_line(self):
        frame, issues, _ = read_docx_questions(
            make_pipe_answer_key_source(), "Pipe_Key.docx"
        )

        self.assertEqual(len(frame), 3)
        self.assertEqual(frame["Correct"].tolist(), ["C", "D", "B"])
        self.assertFalse(issues)

    def test_correct_answers_heading_is_consumed_and_invalid_key_is_flagged(self):
        frame, issues, _ = read_docx_questions(
            make_correct_answers_source_with_invalid_key(), "Invalid_Key.docx"
        )

        self.assertEqual(len(frame), 2)
        self.assertEqual(frame["Correct"].tolist(), ["D", "F"])
        self.assertFalse(bool(frame.iloc[0]["Needs_Review"]))
        self.assertTrue(bool(frame.iloc[1]["Needs_Review"]))
        errors = [item for item in issues if item["Severity"] == "ERROR"]
        self.assertEqual(len(errors), 1)
        self.assertIn("Correct answer must be one of A, B, C, D", errors[0]["Issue"])

    def test_multilevel_unlabelled_options_and_trailing_key_import_as_four_questions(self):
        frame, issues, _ = read_docx_questions(
            make_multilevel_unlabelled_source(), "Equine_Re-sit.docx"
        )

        self.assertEqual(len(frame), 4)
        self.assertEqual(frame["No"].tolist(), ["1", "2", "3", "4"])
        self.assertEqual(frame["Correct"].tolist(), ["D", "C", "B", "B"])
        self.assertEqual(frame["Option_Count"].tolist(), [4, 4, 4, 4])
        self.assertNotIn("Answer Key", frame["Option 5"].tolist())
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])
        self.assertFalse([item for item in issues if "answer key" in item["Issue"].lower()])

    def test_named_questions_and_inline_correct_answer_lines_are_high_confidence(self):
        frame, issues, _ = read_docx_questions(
            make_inline_answer_source(), "Inline_Answers.docx"
        )

        self.assertEqual(frame["No"].tolist(), ["1", "2"])
        self.assertEqual(
            frame["Question"].tolist(),
            [
                "Which synthetic choice is correct for item 1?",
                "Which synthetic choice is correct for item 2?",
            ],
        )
        self.assertEqual(frame["Correct"].tolist(), ["C", "B"])
        self.assertEqual(frame["Needs_Review"].tolist(), [False, False])
        self.assertTrue(
            all(value == "Explicit inline correct-answer line" for value in frame["Answer_Evidence"])
        )
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])

    def test_entire_unique_bold_option_is_high_confidence(self):
        frame, issues, _ = read_docx_questions(
            make_bold_answer_source(), "Bold_Answer.docx"
        )

        self.assertEqual(frame.iloc[0]["Correct"], "B")
        self.assertFalse(bool(frame.iloc[0]["Needs_Review"]))
        self.assertEqual(frame.iloc[0]["Answer_Evidence"], "Entire answer option uniquely bold")
        self.assertFalse([item for item in issues if item["Severity"] == "ERROR"])

    def test_partial_bolding_still_requires_review(self):
        frame, _, _ = read_docx_questions(
            make_bold_answer_source(partial=True), "Partial_Bold_Answer.docx"
        )

        self.assertEqual(frame.iloc[0]["Correct"], "B")
        self.assertTrue(bool(frame.iloc[0]["Needs_Review"]))


class WordAnswerKeyBoundaryRegressionTests(unittest.TestCase):
    """Synthetic coverage for reimporting exams with answer-like stem prefixes."""

    def test_exported_exam_with_answers_round_trips_all_questions(self):
        from app import create_exam_docx

        questions = []
        for index in range(50):
            prefix = ["Describe", "A", "B", "C", "D", "E", "F", "1", "6"][index % 9]
            questions.append({
                "Question": f"{prefix} distinguishing feature of synthetic item {index + 1} is:",
                "Options": [{"letter": label, "text": f"Item {index + 1} option {label}"}
                            for label in "ABCD"],
                "Correct": "ABCD"[index % 4],
            })
        data = create_exam_docx(questions, "Synthetic exam", "A", show_answers=True).getvalue()
        frame, issues, _ = read_docx_questions(data, "Exam_A_WITH_ANSWERS.docx")
        self.assertEqual(frame["No"].tolist(), [str(i) for i in range(1, 51)])
        self.assertEqual(frame["Question"].tolist(), [q["Question"] for q in questions])
        self.assertEqual(frame["Correct"].tolist(), [q["Correct"] for q in questions])
        self.assertEqual(frame["Option_Count"].tolist(), [4] * 50)
        for index, label in enumerate("ABCD", start=1):
            self.assertEqual(frame[f"Option {index}"].tolist(),
                             [q["Options"][index - 1]["text"] for q in questions])
        self.assertTrue(all(value == "Explicit inline correct-answer line"
                            for value in frame["Answer_Evidence"]))
        self.assertFalse(issues)

    def test_answer_key_title_does_not_consume_first_question_starting_with_a(self):
        document = Document()
        document.add_paragraph("Answer Key")
        document.add_paragraph("1. A distinguishing feature of this synthetic object is:")
        for label in "ABCD":
            document.add_paragraph(f"{label}. Synthetic option {label}")
        document.add_paragraph("ANSWER: D")
        data = io.BytesIO()
        document.save(data)
        frame, issues, _ = read_docx_questions(data.getvalue(), "Synthetic.docx")
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["Option_Count"], 4)
        self.assertEqual(frame.iloc[0]["Correct"], "D")
        self.assertFalse(issues)

    def test_numbered_key_requires_complete_pairs_without_question_text(self):
        from question_importer import _numbered_answer_pairs

        for text in ["28. A distinguishing feature is:", "1. B-cell function?",
                     "2. 6 weeks later?", "1. A | 2. B describes a feature"]:
            with self.subTest(text=text):
                self.assertEqual(_numbered_answer_pairs(text), {})
        for separator in [" | ", "; ", ", "]:
            with self.subTest(separator=separator):
                self.assertEqual(_numbered_answer_pairs(separator.join(["1. C", "Q2: 4", "3) B"])),
                                 {1: "C", 2: "D", 3: "B"})

    def test_genuine_eight_option_question_is_still_rejected(self):
        document = Document()
        document.add_paragraph("27. Which synthetic choice is correct?")
        for label in "ABCDEFGH":
            document.add_paragraph(f"{label}. Synthetic option {label}")
        document.add_paragraph("ANSWER: D")
        data = io.BytesIO()
        document.save(data)
        with self.assertRaisesRegex(ValueError, "Question 27 has 8 options"):
            read_docx_questions(data.getvalue(), "Eight_options.docx")


if __name__ == "__main__":
    unittest.main()
