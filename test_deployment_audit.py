import io
import random
import unittest
import zipfile

import pandas as pd
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.shared import Inches
from openpyxl import load_workbook
from PIL import Image
from streamlit.testing.v1 import AppTest

import app
import question_importer
from test_regressions import make_bold_answer_source, make_mcq_32_source


class WordCropStressTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        image = Image.new("RGB", (257, 193), (34, 120, 210))
        output = io.BytesIO()
        image.save(output, format="PNG")
        cls.image_blob = output.getvalue()

    def test_250_random_crop_geometries_have_expected_dimensions(self):
        rng = random.Random(15_011)
        for case_number in range(250):
            left = rng.randint(0, 35_000)
            right = rng.randint(0, 35_000)
            top = rng.randint(0, 35_000)
            bottom = rng.randint(0, 35_000)
            crop = (left, top, right, bottom)
            name, cropped_blob, warning = question_importer._apply_word_image_crop(
                "stress.png", self.image_blob, crop
            )
            expected_size = (
                round(257 * (100_000 - right) / 100_000)
                - round(257 * left / 100_000),
                round(193 * (100_000 - bottom) / 100_000)
                - round(193 * top / 100_000),
            )
            with self.subTest(case=case_number, crop=crop):
                self.assertEqual(name, "stress.png")
                self.assertEqual(warning, "")
                with Image.open(io.BytesIO(cropped_blob)) as cropped:
                    self.assertEqual(cropped.size, expected_size)

    def test_invalid_crop_falls_back_with_visible_warning(self):
        name, result, warning = question_importer._apply_word_image_crop(
            "invalid.png", self.image_blob, (60_000, 0, 40_000, 0)
        )
        self.assertEqual(name, "invalid.png")
        self.assertEqual(result, self.image_blob)
        self.assertIn("Invalid Word crop values", warning)

    def test_drawingml_and_legacy_vml_percentage_forms(self):
        drawingml_cases = {"25000": 25_000, "25%": 25_000, None: 0}
        for value, expected in drawingml_cases.items():
            with self.subTest(kind="DrawingML", value=value):
                self.assertEqual(
                    question_importer._parse_drawingml_crop_value(value), expected
                )

        vml_cases = {
            "16384f": 25_000,
            "25%": 25_000,
            "0.25": 25_000,
            "25000": 25_000,
        }
        for value, expected in vml_cases.items():
            with self.subTest(kind="VML", value=value):
                self.assertEqual(question_importer._parse_vml_crop_value(value), expected)

    def test_same_word_media_with_two_crops_remains_two_visible_images(self):
        source_image = Image.new("RGB", (400, 200), "white")
        for x in range(400):
            colour = (220, 20, 20) if x < 200 else (20, 20, 220)
            for y in range(200):
                source_image.putpixel((x, y), colour)
        image_output = io.BytesIO()
        source_image.save(image_output, format="PNG")

        document = Document()
        document.add_paragraph("32. Which two cropped halves are displayed?")
        image_paragraph = document.add_paragraph()
        crops = [{"r": "50000"}, {"l": "50000"}]
        for crop_values in crops:
            run = image_paragraph.add_run()
            run.add_picture(io.BytesIO(image_output.getvalue()), width=Inches(2))
            blip = run._r.xpath(".//a:blip")[0]
            source_rect = OxmlElement("a:srcRect")
            for side, value in crop_values.items():
                source_rect.set(side, value)
            blip_fill = blip.getparent()
            blip_fill.insert(blip_fill.index(blip) + 1, source_rect)
        for label, text in [("A", "Red"), ("B", "Blue"), ("C", "Both"), ("D", "Neither")]:
            run = document.add_paragraph().add_run(f"{label}. {text}")
            if label == "C":
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        document_output = io.BytesIO()
        document.save(document_output)

        frame, issues, image_map = question_importer.read_docx_questions(
            document_output.getvalue(), "Two_Crops_Question_32.docx"
        )
        image_names = frame.iloc[0]["Image_Files"].split("|")
        self.assertEqual(len(image_names), 2)
        self.assertEqual(len(image_map), 2)
        self.assertTrue(any(item["Severity"] == "WARNING" for item in issues))
        warning = next(item for item in issues if item["Severity"] == "WARNING")
        self.assertEqual(warning["Row"], "32")
        self.assertEqual(warning["Internal_ID"], "Two_Crops_Question_32-32")


class ImportSecurityTests(unittest.TestCase):
    def test_malformed_office_container_is_rejected_cleanly(self):
        with self.assertRaisesRegex(ValueError, "valid modern Office container"):
            question_importer.validate_office_container(b"not-a-zip", "broken.docx")

    def test_member_and_expansion_limits_are_enforced(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("one", b"123456")
            archive.writestr("two", b"abcdef")

        old_member_limit = question_importer.MAX_OFFICE_MEMBERS
        old_expansion_limit = question_importer.MAX_OFFICE_UNCOMPRESSED_BYTES
        try:
            question_importer.MAX_OFFICE_MEMBERS = 1
            with self.assertRaisesRegex(ValueError, "number of internal parts"):
                question_importer.validate_office_container(output.getvalue(), "many.docx")

            question_importer.MAX_OFFICE_MEMBERS = old_member_limit
            question_importer.MAX_OFFICE_UNCOMPRESSED_BYTES = 10
            with self.assertRaisesRegex(ValueError, "1 GB safety limit"):
                question_importer.validate_office_container(output.getvalue(), "large.docx")
        finally:
            question_importer.MAX_OFFICE_MEMBERS = old_member_limit
            question_importer.MAX_OFFICE_UNCOMPRESSED_BYTES = old_expansion_limit

    def test_formula_like_excel_values_are_neutralised(self):
        for value in ["=1+1", "+SUM(A1:A2)", "-2+3", "@cmd", "  =HYPERLINK(\"x\")"]:
            with self.subTest(value=value):
                self.assertTrue(app.safe_excel_value(value).startswith("'"))
        self.assertEqual(app.safe_excel_value("Ordinary text"), "Ordinary text")


class ExamPipelineTests(unittest.TestCase):
    def setUp(self):
        self.frame, self.issues, self.image_map = question_importer.read_docx_questions(
            make_mcq_32_source(), "MCQ_32_Cropped_Image.docx"
        )

    def test_shuffle_preserves_correct_answer_identity_across_100_versions(self):
        for seed in range(100):
            questions = app.build_version_questions(
                self.frame, shuffle_questions=True, shuffle_answers=True, seed=seed
            )
            question = questions[0]
            correct_option = next(
                option for option in question["Options"] if option["letter"] == question["Correct"]
            )
            self.assertEqual(correct_option["text"], "The nodes of Ranvier")

    def test_position_dependent_answer_order_is_not_shuffled(self):
        frame = self.frame.copy()
        frame.loc[0, "Option 4"] = "All of the above"
        for seed in range(20):
            question = app.build_version_questions(
                frame, shuffle_questions=False, shuffle_answers=True, seed=seed
            )[0]
            self.assertTrue(question["Answer_Shuffle_Skipped"])
            self.assertEqual(
                [option["original_letter"] for option in question["Options"]],
                list("ABCD"),
            )

    def test_final_export_package_is_complete_and_keeps_cropped_media(self):
        versions = app.build_versions(
            self.frame,
            num_versions=2,
            naming="Letters",
            shuffle_questions=True,
            shuffle_answers=True,
        )
        report = app.normalise_validation_report(pd.DataFrame())
        package = app.create_zip_output_from_versions(
            versions,
            combined_df=self.frame,
            report_df=report,
            exam_title="MCQ 32 Deployment Audit",
            include_master_with_answers=True,
            image_map=self.image_map,
            formatting=app.get_default_formatting(),
            export_options={
                "word": True,
                "answer_key": True,
                "aiken": True,
                "kahoot": True,
                "export_summary": True,
            },
        )

        with zipfile.ZipFile(package) as archive:
            names = set(archive.namelist())
            for label in ("A", "B"):
                self.assertIn(f"Exam_{label}.docx", names)
                self.assertIn(f"Answer_Key_{label}.docx", names)
                self.assertIn(f"Exam_{label}_WITH_ANSWERS.docx", names)
                self.assertIn(f"Aiken_Blackboard_Version_{label}.txt", names)
                self.assertIn(f"Kahoot_Import_Version_{label}.xlsx", names)

                exam_blob = archive.read(f"Exam_{label}.docx")
                parsed_exam = Document(io.BytesIO(exam_blob))
                exam_text = "\n".join(p.text for p in parsed_exam.paragraphs)
                self.assertIn("longitudinal histologic picture", exam_text)
                with zipfile.ZipFile(io.BytesIO(exam_blob)) as exam_archive:
                    media_name = next(
                        name
                        for name in exam_archive.namelist()
                        if name.startswith("word/media/")
                    )
                    with Image.open(io.BytesIO(exam_archive.read(media_name))) as image:
                        self.assertEqual(image.size, (200, 200))

            mapping = load_workbook(
                io.BytesIO(archive.read("Version_Mapping_Report.xlsx")), data_only=False
            )
            self.assertEqual(mapping["Version Mapping"].max_row, 3)
            validation = load_workbook(
                io.BytesIO(archive.read("Validation_Report.xlsx")), data_only=False
            )
            self.assertIn("Validation Issues", validation.sheetnames)
            self.assertIn("Imported Questions", validation.sheetnames)
            self.assertEqual(archive.testzip(), None)

    def test_package_builder_uses_the_explicit_blueprint(self):
        versions = app.build_versions(
            self.frame,
            num_versions=1,
            naming="Letters",
            shuffle_questions=False,
            shuffle_answers=False,
        )
        report = app.normalise_validation_report(pd.DataFrame())
        blueprint = app.make_blueprint_dataframe(
            self.frame,
            {self.frame.iloc[0]["Source_File"]: 1},
        )

        package = app.create_zip_output_from_versions(
            versions,
            combined_df=self.frame,
            report_df=report,
            exam_title="Blueprint audit",
            include_master_with_answers=False,
            image_map=self.image_map,
            formatting=app.get_default_formatting(),
            blueprint_df=blueprint,
        )

        with zipfile.ZipFile(package) as archive:
            self.assertIn("Blueprint_Report.xlsx", archive.namelist())
            report_book = load_workbook(
                io.BytesIO(archive.read("Blueprint_Report.xlsx")), data_only=True
            )
            worksheet = report_book["Blueprint"]
            self.assertEqual(worksheet["A2"].value, self.frame.iloc[0]["Source_File"])
            self.assertEqual(worksheet["C2"].value, 1)


class StreamlitRuntimeTests(unittest.TestCase):
    def test_changed_settings_remove_stale_package_and_download_navigation(self):
        app_test = AppTest.from_file("app.py", default_timeout=60).run()
        app_test.get("file_uploader")[0].upload(
            "sample.txt", b"1. Which number is even?\nA. Two\nB. Three\nANSWER: A\n", "text/plain"
        ).run()
        next(button for button in app_test.button if button.label == "Create exam package").click().run()
        self.assertTrue(app_test.session_state["preview_versions"])
        next(field for field in app_test.text_input if field.label == "Exam title").set_value("Updated exam").run()
        self.assertEqual(len(app_test.exception), 0)
        self.assertIsNone(app_test.session_state["preview_versions"])
        self.assertFalse(any("Download exam package" in item.label for item in app_test.get("download_button")))
        navigation = [item.value for item in app_test.markdown if '<nav class="workflow-shell"' in item.value]
        self.assertEqual(len(navigation), 1)
        self.assertNotIn('href="#download-package"', navigation[0])

    def test_export_controls_preserve_configuration_across_reruns(self):
        app_test = AppTest.from_file("app.py", default_timeout=60).run()
        app_test.get("file_uploader")[0].upload(
            "sample.txt", b"1. Which number is even?\nA. Two\nB. Three\nANSWER: A\n", "text/plain"
        ).run()
        next(field for field in app_test.number_input if field.label == "Number of exam versions").set_value(3).run()
        next(field for field in app_test.checkbox if field.label == "Export Kahoot import Excel").check().run()
        next(field for field in app_test.selectbox if field.label == "Kahoot time limit (seconds)").select(60).run()
        next(field for field in app_test.checkbox if field.label == "Include lecturer copies with answers").uncheck().run()
        next(button for button in app_test.button if button.label == "Create exam package").click().run()
        self.assertEqual(len(app_test.exception), 0)
        settings = app_test.session_state["preview_settings"]
        self.assertEqual(len(app_test.session_state["preview_versions"]), 3)
        self.assertTrue(settings["export_options"]["kahoot"])
        self.assertEqual(settings["kahoot_time_limit"], 60)
        self.assertFalse(settings["include_master_with_answers"])

    def test_cold_start_has_no_app_exception(self):
        app_test = AppTest.from_file("app.py", default_timeout=30).run()
        self.assertEqual(len(app_test.exception), 0)
        self.assertEqual(app_test.title[0].value, "MCQ Exam Builder")
        about = next(item for item in app_test.expander if item.label == "About")
        about_text = "\n".join(item.value for item in about.markdown)
        self.assertIn("Christos I. Karagiannis", about_text)
        self.assertIn("does not write questions or judge academic correctness", about_text)
        privacy = next(item for item in app_test.expander if item.label == "Privacy & data handling")
        privacy_text = "\n".join(item.value for item in privacy.markdown)
        self.assertIn("not solely in your browser", privacy_text)
        self.assertIn("not a guarantee of immediate deletion", privacy_text)

    def test_mcq_32_upload_preview_and_download_path(self):
        app_test = AppTest.from_file("app.py", default_timeout=60).run()
        app_test.get("file_uploader")[0].upload(
            "MCQ_32_Cropped_Image.docx",
            make_mcq_32_source(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ).run()
        self.assertEqual(len(app_test.exception), 0)
        image_blob = next(iter(app_test.session_state["image_map"].values()))
        with Image.open(io.BytesIO(image_blob)) as image:
            self.assertEqual(image.size, (200, 200))

        next(
            button
            for button in app_test.button
            if button.label == "Preview exam first"
        ).click().run()
        self.assertEqual(len(app_test.exception), 0)
        self.assertIn("A", app_test.session_state["preview_versions"])
        self.assertIn("Exam preview", [item.value for item in app_test.subheader])
        self.assertEqual(
            [header.value.split(" — ")[0] for header in app_test.header],
            ["1", "2", "3", "4"],
        )
        self.assertTrue(
            any("Download exam package" in button.label for button in app_test.get("download_button"))
        )

    def test_direct_download_skips_preview_and_creates_package(self):
        app_test = AppTest.from_file("app.py", default_timeout=60).run()
        app_test.get("file_uploader")[0].upload(
            "MCQ_32_Cropped_Image.docx",
            make_mcq_32_source(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ).run()

        next(
            button
            for button in app_test.button
            if button.label == "Create exam package"
        ).click().run()

        self.assertEqual(len(app_test.exception), 0)
        self.assertFalse(app_test.session_state["preview_settings"]["review_before_download"])
        self.assertEqual(
            [header.value.split(" — ")[0] for header in app_test.header],
            ["1", "2", "3", "4"],
        )
        self.assertNotIn(
            "Exam preview",
            [item.value for item in app_test.subheader],
        )
        self.assertIn(
            "4 — Download",
            [header.value for header in app_test.header],
        )
        self.assertTrue(
            any("Download exam package" in button.label for button in app_test.get("download_button"))
        )


    def test_warning_banner_and_location_identify_question_32(self):
        warning_source = b"""32. Which choice depends on its displayed position?
A. First
B. Second
C. Third
D. All of the above
ANSWER: D
"""
        app_test = AppTest.from_file("app.py", default_timeout=30).run()
        app_test.get("file_uploader")[0].upload(
            "Warning_Location_32.txt", warning_source, "text/plain"
        ).run()
        self.assertEqual(len(app_test.exception), 0)
        self.assertTrue(
            any("Validation details" in warning.value for warning in app_test.warning)
        )
        self.assertTrue(
            any("1 warning needs your attention" in warning.value for warning in app_test.warning)
        )
        self.assertIn(
            "Validation details — affected questions and locations",
            [subheader.value for subheader in app_test.subheader],
        )
        report = app_test.session_state["report_df"]
        self.assertEqual(len(report), 1)
        self.assertIn("Warning_Location_32.txt", report.iloc[0]["Location"])
        self.assertIn("ID Warning_Location_32-32", report.iloc[0]["Location"])


    def test_required_review_gate_is_prominent_and_preview_explains_why_locked(self):
        app_test = AppTest.from_file("app.py", default_timeout=30).run()
        app_test.get("file_uploader")[0].upload(
            "Partial_Bold_Answer.docx",
            make_bold_answer_source(partial=True),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ).run()

        self.assertEqual(len(app_test.exception), 0)
        self.assertTrue(
            any("Action required" in warning.value for warning in app_test.warning)
        )
        self.assertIn(
            "Required action — confirm flagged questions",
            [subheader.value for subheader in app_test.subheader],
        )
        locked_buttons = [
            button
            for button in app_test.button
            if button.label in {
                "Create exam package — review required",
                "Preview exam first — review required",
            }
        ]
        self.assertEqual(len(locked_buttons), 2)
        self.assertTrue(all(button.disabled for button in locked_buttons))


if __name__ == "__main__":
    unittest.main()
